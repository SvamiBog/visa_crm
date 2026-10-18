"""Reading an uploaded spreadsheet (backlog 1.14): XLSX — a Google Sheets download with all
its sheets — or CSV. The file comes from the user, so it is bounded before it is read:
its size, the unpacked size of an XLSX (a zip), the number of rows and columns. XML is
parsed with defusedxml (openpyxl picks it up when installed).
"""

import csv
import io
import itertools
import re
import warnings
import zipfile
from collections import Counter
from dataclasses import dataclass

from app.core.import_values import Cell, cell_text, is_blank

MAX_BYTES = 5 * 1024 * 1024
MAX_UNPACKED = 60 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 100
# The header is looked for in the first rows: a sheet may start with a title.
HEADER_SEARCH = 20


class FileError(ValueError):
    """A file that cannot be imported; the code is the API error (errors.* in locales)."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class SheetRow:
    # The row number as the specialist sees it in the spreadsheet.
    number: int
    cells: tuple[Cell, ...]


@dataclass(frozen=True)
class Table:
    sheets: list[str]
    sheet: str
    header_row: int
    headers: list[str]
    rows: list[SheetRow]


def _xlsx_rows(data: bytes, sheet: str | None) -> tuple[list[str], str, list[tuple[Cell, ...]]]:
    import openpyxl  # imported here: only uploads need it

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED:
                raise FileError("file_too_large")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except FileError:
        raise
    except Exception as error:  # a broken or foreign file: openpyxl raises many kinds
        raise FileError("file_unreadable") from error
    try:
        names = list(book.sheetnames)
        name = sheet if sheet is not None else names[0]
        if name not in names:
            raise FileError("sheet_not_found")
        rows: list[tuple[Cell, ...]] = []
        for row in book[name].iter_rows(values_only=True):
            rows.append(tuple(row[:MAX_COLUMNS]))
            if len(rows) > MAX_ROWS + HEADER_SEARCH + 1000:
                raise FileError("too_many_rows")
        return names, name, rows
    finally:
        book.close()


def _csv_rows(data: bytes) -> list[tuple[Cell, ...]]:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise FileError("file_unreadable")
    if "\x00" in text:
        raise FileError("file_unreadable")
    reader = csv.reader(io.StringIO(text), delimiter=_delimiter(text))
    rows: list[tuple[Cell, ...]] = []
    for row in reader:
        rows.append(tuple(row[:MAX_COLUMNS]))
        if len(rows) > MAX_ROWS + HEADER_SEARCH + 1000:
            raise FileError("too_many_rows")
    return rows


def _delimiter(text: str) -> str:
    """Google Sheets writes commas, Excel with Russian settings — semicolons. The one that
    splits the most of the first rows into the same number of columns wins."""
    best, best_score = ",", (0, 0)
    for delimiter in (",", ";", "\t"):
        widths = Counter(
            len(row)
            for row in itertools.islice(csv.reader(io.StringIO(text), delimiter=delimiter), 50)
            if len(row) > 1
        )
        if widths:
            width, rows = widths.most_common(1)[0]
            if (rows, width) > best_score:
                best, best_score = delimiter, (rows, width)
    return best


_ZERO = re.compile(r"^0+(?:[.,]0+)?$")


def _filled(row: tuple[Cell, ...]) -> int:
    """Cells with something in them: the zeros and dashes of a sheet's placeholder rows
    after the last client do not count."""
    return sum(
        1
        for cell in row
        if (text := cell_text(cell)) is not None and not is_blank(text) and not _ZERO.match(text)
    )


def read_table(data: bytes, sheet: str | None = None, header_row: int | None = None) -> Table:
    """The sheet's header and data rows. The header is the given row or the first one with
    at least two filled cells; empty rows — and rows of only zeros and dashes — are
    dropped, row numbers are kept."""
    if not data:
        raise FileError("file_empty")
    if len(data) > MAX_BYTES:
        raise FileError("file_too_large")
    if data.startswith(b"PK\x03\x04"):
        sheets, name, rows = _xlsx_rows(data, sheet)
    elif data.startswith(b"\xd0\xcf\x11\xe0"):  # the old binary .xls
        raise FileError("xls_not_supported")
    else:
        rows = _csv_rows(data)
        sheets, name = [""], ""
        if sheet:
            raise FileError("sheet_not_found")

    if header_row is None:
        found = next((i for i, row in enumerate(rows[:HEADER_SEARCH]) if _filled(row) >= 2), None)
        if found is None:
            raise FileError("no_header")
        header_index = found
    elif 1 <= header_row <= min(len(rows), HEADER_SEARCH):
        header_index = header_row - 1
    else:
        raise FileError("no_header")

    body = [
        SheetRow(number, row)
        for number, row in enumerate(rows[header_index + 1 :], start=header_index + 2)
        if _filled(row)
    ]
    if len(body) > MAX_ROWS:
        raise FileError("too_many_rows")
    width = max([_last_filled(rows[header_index])] + [_last_filled(r.cells) for r in body])
    headers = [cell_text(cell) or "" for cell in rows[header_index]]
    headers += [""] * (width - len(headers))
    headers = headers[:width]
    body = [SheetRow(r.number, (r.cells + (None,) * width)[:width]) for r in body]
    return Table(sheets, name, header_index + 1, headers, body)


def _last_filled(cells: tuple[Cell, ...]) -> int:
    for i in range(len(cells), 0, -1):
        if cell_text(cells[i - 1]) is not None:
            return i
    return 0
