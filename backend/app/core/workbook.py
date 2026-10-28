"""Reading a file in the VisaCRM data format (backlog 1.21, docs/data-format.md): one XLSX
workbook, every sheet at once. The file comes from the user, so it is bounded before it is
read — its size, its unpacked size (an XLSX is a zip), the rows and columns. XML is parsed
with defusedxml (openpyxl picks it up when installed).

Sheets and headers are matched by their Russian or English names; anything else is passed
by with a warning. Each cell is read strictly (app/core/data_values.py). A bad cell in a
required column skips its row, in another column only the cell; the report names the
sheet, the row and the column, never the value.
"""

import io
import re
import warnings
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from app.core.data_format import (
    ABOUT_EXPORTED,
    ABOUT_SHEET,
    CODE_LENGTH,
    REFERENCE_SHEET,
    SHEETS,
    Column,
    ColumnKind,
    Sheet,
    ValueType,
)
from app.core.data_values import (
    Cell,
    CellError,
    cell_text,
    normalize,
    parse_bool,
    parse_choice,
    parse_currency,
    parse_date,
    parse_document_number,
    parse_email,
    parse_int,
    parse_list,
    parse_money,
    parse_percent,
    parse_phone,
    parse_rate,
    parse_telegram,
    parse_text,
)

MAX_BYTES = 10 * 1024 * 1024
MAX_UNPACKED = 120 * 1024 * 1024
# Rows of data on all sheets together.
MAX_ROWS = 20_000
MAX_COLUMNS = 100
# Empty rows are passed by, but a sheet is not read without end.
MAX_EMPTY_ROWS = 5000
# A name or a list value: looked up in the CRM, so its length is checked there.
LOOKUP_LENGTH = 200
SECRET_LENGTH = 200
_EXPORTED_TEXT = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})\s+(\d{1,2}):(\d{2})$")


class WorkbookError(ValueError):
    """A file that cannot be imported; the code is the API error (errors.* in locales)."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class Issue:
    """A line of the report: where and why, without the cell's value."""

    reason: str
    sheet: str
    row: int | None = None
    column: str | None = None


@dataclass
class Row:
    # The row number as the specialist sees it in the file.
    number: int
    # Parsed values of the filled cells, by column key.
    values: dict[str, Any] = field(default_factory=dict)
    # Columns present on the sheet whose cell is empty: a snapshot row clears them.
    blank: set[str] = field(default_factory=set)


@dataclass
class SheetData:
    spec: Sheet
    # The sheet's name as in the file.
    title: str
    # Column key → its header as in the file, for the report.
    headers: dict[str, str]
    rows: list[Row] = field(default_factory=list)


@dataclass
class Workbook:
    sheets: dict[str, SheetData] = field(default_factory=dict)
    # «О файле» → «Выгружено»: the file is a CRM export (a snapshot), naive local time.
    exported_at: datetime | None = None
    skipped: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)


_BY_NAME = {normalize(name): sheet for sheet in SHEETS for name in (sheet.ru, sheet.en)}
_RUSSIAN = {normalize(sheet.ru) for sheet in SHEETS}
_PASSED_BY = {normalize(name) for name in REFERENCE_SHEET}
_ABOUT = {normalize(name) for name in ABOUT_SHEET}
_EXPORTED = {normalize(name) for name in ABOUT_EXPORTED}


def read_workbook(data: bytes) -> Workbook:
    """Every sheet of the format in the file, parsed; raises WorkbookError for a file that
    cannot be read at all."""
    if len(data) > MAX_BYTES:
        raise WorkbookError("file_too_large")
    result = Workbook()
    counter = _RowCounter()
    for title, rows in _xlsx_sheets(data):
        name = normalize(title)
        if name in _ABOUT:
            result.exported_at = _exported_at(rows)
            continue
        if name in _PASSED_BY:
            continue
        spec = _BY_NAME.get(name)
        if spec is None:
            result.warnings.append(Issue("unknown_sheet", title))
            continue
        if spec.key in result.sheets:
            result.warnings.append(Issue("duplicate_sheet", title))
            continue
        sheet = _read_sheet(spec, title, name in _RUSSIAN, rows, result, counter)
        if sheet is not None:
            result.sheets[spec.key] = sheet
    return result


class _RowCounter:
    def __init__(self) -> None:
        self.rows = 0

    def add(self) -> None:
        self.rows += 1
        if self.rows > MAX_ROWS:
            raise WorkbookError("too_many_rows")


def _xlsx_sheets(data: bytes) -> Iterator[tuple[str, Iterator[tuple[Cell, ...]]]]:
    import openpyxl  # imported here: only uploads need it

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED:
                raise WorkbookError("file_too_large")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except WorkbookError:
        raise
    except Exception as error:  # a broken or foreign file: openpyxl raises many kinds
        raise WorkbookError("file_unreadable") from error
    try:
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            yield sheet.title, (tuple(row[:MAX_COLUMNS]) for row in rows)
    finally:
        book.close()


def _header(column: Column, russian: bool) -> str:
    return column.ru if russian else column.en


def _read_sheet(
    spec: Sheet,
    title: str,
    russian: bool,
    rows: Iterator[tuple[Cell, ...]],
    result: Workbook,
    counter: _RowCounter,
) -> SheetData | None:
    header_row = next(rows, None)
    if header_row is None:
        return None
    by_header = {normalize(name): c for c in spec.columns for name in (c.ru, c.en)}
    places: dict[str, int] = {}
    headers: dict[str, str] = {}
    for index, cell in enumerate(header_row):
        text = cell_text(cell)
        if text is None:
            continue
        column = by_header.get(normalize(text))
        if column is None:
            result.warnings.append(Issue("unknown_column", title, column=text))
        elif column.key in places:
            result.warnings.append(Issue("duplicate_column", title, column=text))
        else:
            places[column.key] = index
            headers[column.key] = text
    missing = [c for c in spec.columns if c.required and c.key not in places]
    if missing:
        for column in missing:
            result.skipped.append(Issue("missing_column", title, column=_header(column, russian)))
        return None

    sheet = SheetData(spec, title, headers)
    codes: set[str] = set()
    empty = 0
    for number, cells in enumerate(rows, start=2):
        if all(cell_text(cell) is None for cell in cells):
            empty += 1
            if empty > MAX_EMPTY_ROWS:
                break
            continue
        empty = 0
        counter.add()
        row = _read_row(spec, title, number, cells, places, headers, result)
        if row is None:
            continue
        code = row.values.get("code")
        if code is not None:
            if code in codes:
                result.skipped.append(
                    Issue("duplicate_code", title, number, headers[spec.code.key])
                )
                continue
            codes.add(code)
        sheet.rows.append(row)
    return sheet


def _read_row(
    spec: Sheet,
    title: str,
    number: int,
    cells: tuple[Cell, ...],
    places: dict[str, int],
    headers: dict[str, str],
    result: Workbook,
) -> Row | None:
    row = Row(number)
    for column in spec.columns:
        if column.key not in places:
            continue
        index = places[column.key]
        cell = cells[index] if index < len(cells) else None
        if column.kind == ColumnKind.REF:
            if column.key == "client_name" and (text := cell_text(cell)):
                row.values["client_name"] = " ".join(text.split())
            continue
        if cell_text(cell) is None:
            if column.required:
                result.skipped.append(Issue("required", title, number, headers[column.key]))
                return None
            row.blank.add(column.key)
            continue
        try:
            row.values[column.key] = _parse(column, cell)
        except CellError as error:
            issue = Issue(error.reason, title, number, headers[column.key])
            if column.required:
                result.skipped.append(issue)
                return None
            result.warnings.append(issue)
    return row


def _parse(column: Column, cell: Cell) -> Any:
    match column.type:
        case ValueType.CODE:
            return parse_text(cell, column.max_length or CODE_LENGTH)
        case ValueType.TEXT:
            return parse_text(cell, column.max_length)
        case ValueType.DOCUMENT_NUMBER:
            return parse_document_number(cell)
        case ValueType.DATE:
            return parse_date(cell)
        case ValueType.INT:
            return parse_int(cell)
        case ValueType.MONEY:
            return parse_money(cell)
        case ValueType.RATE:
            return parse_rate(cell)
        case ValueType.PERCENT:
            return parse_percent(cell)
        case ValueType.CURRENCY:
            return parse_currency(cell)
        case ValueType.BOOL:
            return parse_bool(cell)
        case ValueType.PHONE:
            return parse_phone(cell)
        case ValueType.TELEGRAM:
            return parse_telegram(cell)
        case ValueType.EMAIL:
            return parse_email(cell)
        case ValueType.CHOICE:
            return parse_choice(cell, column.choices)
        case ValueType.COUNTRIES:
            return parse_list(cell)
        case ValueType.SECRET:
            # As typed: a password is not tidied up.
            text = cell if isinstance(cell, str) else cell_text(cell) or ""
            if len(text) > (column.max_length or SECRET_LENGTH):
                raise CellError("too_long")
            return text
    # Countries, cities and the specialist's lists: looked up by the import.
    return parse_text(cell, LOOKUP_LENGTH)


def _exported_at(rows: Iterator[tuple[Cell, ...]]) -> datetime | None:
    for cells in rows:
        if len(cells) < 2 or normalize(cell_text(cells[0]) or "") not in _EXPORTED:
            continue
        value = cells[1]
        if isinstance(value, datetime):
            return value
        if isinstance(value, str) and (match := _EXPORTED_TEXT.match(value.strip())):
            day, month, year, hour, minute = (int(g) for g in match.groups())
            try:
                return datetime(year, month, day, hour, minute)
            except ValueError:
                return None
    return None
