"""Files in the VisaCRM data format the CRM gives (backlog 1.21, docs/data-format.md →
«Выгрузка»): the template — every sheet with its headers, the «Справочник» with the values
the workspace's lists and the format's fixed words allow, drop-downs from it, and «О файле».
The specialist gives the template to the AI with the prompt (docs/import-prompt.md).
Reference columns are grey, codes in grey letters, the first row stays in place.
"""

import io
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.data_format import (
    ABOUT_SHEET,
    REFERENCE_SHEET,
    SHEETS,
    VERSION,
    Choice,
    Column,
    ColumnKind,
    ValueType,
)
from app.core.money import CURRENCIES

# Rows a drop-down covers below the header.
LIST_ROWS = 5000
HEADER_FONT = Font(bold=True)
KEY_FONT = Font(bold=True, color="808080")
REFERENCE_FILL = PatternFill("solid", fgColor="E7E6E6")
YES_NO = (Choice("yes", "да", "yes"), Choice("no", "нет", "no"))


@dataclass
class Lists:
    """The workspace's values for the «Справочник»."""

    countries: list[str] = field(default_factory=list)
    cities: list[str] = field(default_factory=list)
    channels: list[str] = field(default_factory=list)
    partners: list[str] = field(default_factory=list)
    purposes: list[str] = field(default_factory=list)
    registrars: list[str] = field(default_factory=list)
    systems: list[str] = field(default_factory=list)


_LIST_OF_TYPE = {
    ValueType.COUNTRY: "countries",
    ValueType.CITY: "cities",
    ValueType.CHANNEL: "channels",
    ValueType.PARTNER: "partners",
    ValueType.PURPOSE: "purposes",
    ValueType.REGISTRAR: "registrars",
    ValueType.SYSTEM: "systems",
}
_LIST_TITLES = {
    "countries": ("Страны", "Countries"),
    "cities": ("Города подачи", "Application cities"),
    "channels": ("Каналы прихода", "Lead channels"),
    "partners": ("Партнёры", "Partners"),
    "purposes": ("Цели", "Purposes"),
    "registrars": ("Регистраторы", "Registrars"),
    "systems": ("Системы сейфа", "Vault systems"),
    "currencies": ("Валюты", "Currencies"),
    "yes_no": ("Да / нет", "Yes / no"),
}
_ABOUT_ROWS = (
    (("Формат", "Format"), (f"VisaCRM, версия {VERSION}", f"VisaCRM, version {VERSION}")),
)


def _word(language: str, ru: str, en: str) -> str:
    return ru if language == "ru" else en


def _values(column: Column, lists: Lists, language: str) -> tuple[str, list[str]] | None:
    """The «Справочник» column a column's drop-down takes its values from."""
    if column.kind != ColumnKind.DATA:
        return None
    if column.type == ValueType.CHOICE:
        return f"{column.key}", [_word(language, c.ru, c.en) for c in column.choices]
    if column.type == ValueType.BOOL:
        return "yes_no", [_word(language, c.ru, c.en) for c in YES_NO]
    if column.type == ValueType.CURRENCY:
        return "currencies", list(CURRENCIES)
    if (name := _LIST_OF_TYPE.get(column.type)) is not None:
        return name, getattr(lists, name)
    return None


def template(lists: Lists, *, language: str, workspace: str, base_currency: str) -> bytes:
    """The empty workbook of the format, in the interface's language."""
    book = Workbook()
    book.remove(book.active)
    data_sheets = [
        (sheet, book.create_sheet(_word(language, sheet.ru, sheet.en))) for sheet in SHEETS
    ]
    lists_sheet = book.create_sheet(_word(language, *REFERENCE_SHEET))
    ranges: dict[str, str] = {}

    def list_range(name: str, title: str, values: list[str]) -> str | None:
        if not values:
            return None
        if name not in ranges:
            col = len(ranges) + 1
            letter = get_column_letter(col)
            head = lists_sheet.cell(1, col, title)
            head.font, head.fill = HEADER_FONT, REFERENCE_FILL
            for row, value in enumerate(values, start=2):
                lists_sheet.cell(row, col, value)
            lists_sheet.column_dimensions[letter].width = max(16, len(title) + 2)
            quoted = lists_sheet.title.replace("'", "''")
            ranges[name] = f"'{quoted}'!${letter}$2:${letter}${len(values) + 1}"
        return ranges[name]

    for spec, sheet in data_sheets:
        for col, column in enumerate(spec.columns, start=1):
            letter = get_column_letter(col)
            cell = sheet.cell(1, col, _word(language, column.ru, column.en))
            cell.font = KEY_FONT if column.kind == ColumnKind.KEY else HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if column.kind == ColumnKind.REF:
                cell.fill = REFERENCE_FILL
            sheet.column_dimensions[letter].width = max(10, min(28, len(cell.value) + 2))
            found = _values(column, lists, language)
            if found is None:
                continue
            name, values = found
            title = _LIST_TITLES.get(name, (column.ru, column.en))
            target = list_range(name, _word(language, *title), values)
            if target is None:
                continue
            check = DataValidation(type="list", formula1=f"={target}", allow_blank=True)
            check.showErrorMessage = column.type in (ValueType.CHOICE, ValueType.BOOL)
            sheet.add_data_validation(check)
            check.add(f"{letter}2:{letter}{LIST_ROWS}")
        sheet.row_dimensions[1].height = 32
        sheet.freeze_panes = "B2"
    lists_sheet.freeze_panes = "A2"

    about = book.create_sheet(_word(language, *ABOUT_SHEET))
    rows = [
        *((_word(language, *k), _word(language, *v)) for k, v in _ABOUT_ROWS),
        (_word(language, "Кабинет", "Workspace"), workspace),
        (_word(language, "Базовая валюта", "Base currency"), base_currency),
    ]
    for number, (key, value) in enumerate(rows, start=1):
        about.cell(number, 1, key).font = HEADER_FONT
        about.cell(number, 2, value)
    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 48

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
