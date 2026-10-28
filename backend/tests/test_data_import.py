"""Import of the VisaCRM data format (backlog 1.20, docs/data-format.md): one workbook with
every sheet at once, linked by codes, read strictly. A repeated or re-made file finds the
records it made; a row of a CRM export is a snapshot of its record and is not loaded over
a later change; codes of another workspace find nothing; the report never holds a value
of the file."""

import base64
import io
import uuid
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import openpyxl
from httpx import AsyncClient
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine

from tests.helpers import MakeUser, login

API = "/api/imports"


async def ok(response) -> dict:
    assert response.status_code in (200, 201), response.text
    return response.json()


def xlsx(sheets: dict[str, list[list]]) -> str:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode()


def body(sheets: dict[str, list[list]], **extra) -> dict:
    return {"file_name": "VisaCRM-импорт.xlsx", "content": xlsx(sheets)} | extra


async def check(client: AsyncClient, sheets: dict[str, list[list]], **extra) -> dict:
    return await ok(await client.post(f"{API}/check", json=body(sheets, **extra)))


async def run(client: AsyncClient, sheets: dict[str, list[list]], **extra) -> dict:
    return await ok(await client.post(API, json=body(sheets, **extra)))


async def count(engine: AsyncEngine, table: str, tenant_id: uuid.UUID) -> int:
    async with engine.connect() as conn:
        return await conn.scalar(
            text(f"SELECT count(*) FROM {table} WHERE tenant_id = :t"), {"t": tenant_id}
        )


async def rows(engine: AsyncEngine, sql: str, **params) -> list[dict]:
    async with engine.connect() as conn:
        return [dict(r) for r in (await conn.execute(text(sql), params)).mappings()]


async def mail_system(client: AsyncClient) -> None:
    for name, kind in (("Почта", "email"), ("Booking.com", "booking")):
        await ok(
            await client.post("/api/dictionaries/vault-systems", json={"name": name, "kind": kind})
        )


CLIENTS = [
    [
        "Код клиента",
        "Фамилия",
        "Имя",
        "Отчество",
        "Телефон",
        "Telegram",
        "Дата рождения",
        "Канал прихода",
    ],
    [
        "Подача25-2",
        "Петрова",
        "Анна",
        "Сергеевна",
        "+375291112233",
        "@anna_p",
        date(1988, 3, 12),
        "Сарафан",
    ],
    ["Подача25-3", "Петров", "Игорь", None, "+375297778899", None, None, "Сарафан"],
    ["Франция-2", "Сидоренко", "Ольга", None, None, None, None, None],
]
NOTES = [
    ["Код клиента", "Клиент", "Текст"],
    ["Подача25-2", "Петрова Анна Сергеевна", "Звонить после 18:00"],
]
CONSULTATIONS = [
    [
        "Код консультации",
        "Код клиента",
        "Клиент",
        "Дата",
        "Заявителей",
        "Страны",
        "Решение",
        "Стоимость",
        "Валюта",
        "Оплачено",
    ],
    [
        "Конс-2",
        "Подача25-2",
        "Петрова Анна Сергеевна",
        date(2025, 1, 15),
        2,
        "Польша, Латвия",
        "думает",
        50,
        "BYN",
        "да",
    ],
]
CASES = [
    [
        "Код кейса",
        "Код клиента",
        "Клиент",
        "Семья",
        "Код консультации",
        "Источник",
        "Страна визы",
        "Регистратор",
        "Номер в ВЦ",
        "Дата подачи",
        "Результат",
        "Кратность",
        "Дата решения",
        "Стоимость",
        "Талон",
        "Валюта талона",
        "Талон оплачен до CRM",
    ],
    [
        "Подача25-2",
        "Подача25-2",
        "Петрова Анна Сергеевна",
        "Петровы",
        "Конс-2",
        "Подача25",
        "Польша",
        "Иван",
        "PL123",
        date(2025, 5, 20),
        "виза",
        "многократная",
        date(2025, 6, 2),
        800,
        300,
        "PLN",
        "да",
    ],
    [
        "Подача25-3",
        "Подача25-3",
        "Петров Игорь",
        "Петровы",
        "Конс-2",
        "Подача25",
        "Польша",
        "Иван",
        None,
        date(2025, 5, 20),
        None,
        None,
        None,
        600,
        None,
        None,
        None,
    ],
    [
        "Франция-2",
        "Франция-2",
        "Сидоренко Ольга",
        None,
        None,
        "Франция",
        "Франция",
        None,
        None,
        None,
        None,
        None,
        None,
        500,
        None,
        None,
        None,
    ],
]
PAYMENTS = [
    ["Код кейса", "Клиент", "Дата", "Сумма", "Валюта", "Официальный"],
    ["Подача25-2", "Петрова Анна Сергеевна", date(2025, 1, 15), 200, "BYN", "нет"],
    ["Подача25-2", "Петрова Анна Сергеевна", date(2025, 5, 20), 150, "EUR", "да"],
]
EXPENSES = [
    ["Код кейса", "Дата", "Описание", "Сумма"],
    ["Подача25-2", date(2025, 5, 12), "Перевод справки", 30],
]
QUEUE = [
    ["Код кейса", "Клиент", "Дата постановки", "Дата письма"],
    ["Подача25-2", "Петрова Анна Сергеевна", date(2025, 1, 20), date(2025, 5, 5)],
    ["Подача25-3", "Петров Игорь", date(2025, 1, 20), date(2025, 5, 5)],
]
ATTEMPTS = [
    ["Код кейса", "Номер попытки", "Дата отправки", "Последний день окна", "Итог", "Дата закрытия"],
    ["Франция-2", 1, date(2025, 2, 10), date(2025, 4, 11), "истекла", date(2025, 4, 11)],
    ["Франция-2", 2, date(2025, 4, 12), date(2025, 6, 11), None, None],
]
ACCOUNTS = [
    ["Код клиента", "Код кейса", "Система", "Логин", "Пароль", "Чья почта", "Создала сама"],
    ["Подача25-2", "Подача25-2", "Почта", "anna.p@example.com", "Pa55word", "клиента", "нет"],
    [None, None, "Booking.com", "booking.mine@example.com", "BkPass1", "моя резервная", "да"],
]
BOOKINGS = [
    ["Код кейса", "Место", "Заезд", "Выезд", "Бесплатная отмена до", "Номер брони", "Аккаунт"],
    [
        "Подача25-2",
        "Hotel Old Town, Kraków",
        date(2025, 6, 12),
        date(2025, 6, 19),
        date(2025, 6, 5),
        "4455667788",
        "booking.mine@example.com",
    ],
]
BOOK = {
    "Клиенты": CLIENTS,
    "Заметки": NOTES,
    "Консультации": CONSULTATIONS,
    "Кейсы": CASES,
    "Платежи": PAYMENTS,
    "Расходы": EXPENSES,
    "Очередь": QUEUE,
    "Заявки с окном": ATTEMPTS,
    "Учётные записи": ACCOUNTS,
    "Брони": BOOKINGS,
}


async def test_a_workbook_becomes_clients_cases_and_all_their_records(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")
    await mail_system(client)

    report = await run(client, BOOK)

    assert report["skipped_count"] == 0, report["skipped"]
    assert report["warnings_count"] == 0, report["warnings"]
    assert report["created"] == {
        "contacts": 3,
        "notes": 1,
        "consultations": 1,
        "cases": 3,
        "groups": 1,
        "finances": 3,
        "payments": 2,
        "expenses": 1,
        "queue_entries": 2,
        "attempts": 2,
        "vault_entries": 2,
        "passwords": 2,
        "bookings": 1,
        "channel": 1,
        "registrar": 1,
    }
    cases = await rows(
        owner_engine,
        "SELECT c.full_name, v.number, v.group_id, v.stage, v.outcome, v.visa_entries, "
        "v.consultation_id IS NOT NULL AS from_consultation, f.price, f.registrar_fee_amount, "
        "f.registrar_fee_currency, f.registrar_fee_base, f.registrar_fee_settled "
        "FROM visa_cases v JOIN contacts c ON c.id = v.contact_id "
        "LEFT JOIN case_finances f ON f.case_id = v.id WHERE v.tenant_id = :t ORDER BY v.number",
        t=tenant,
    )
    anna, igor, olga = cases
    assert anna["group_id"] is not None and anna["group_id"] == igor["group_id"]
    assert olga["group_id"] is None
    assert (anna["stage"], anna["outcome"], anna["visa_entries"]) == ("result", "visa", "multiple")
    assert igor["stage"] == "review" and olga["stage"] == "waiting"
    assert anna["from_consultation"] and igor["from_consultation"]
    # The talon in złoty at the National Bank rate, already paid to the registrar.
    assert (anna["registrar_fee_amount"], anna["registrar_fee_currency"]) == (300, "PLN")
    assert anna["registrar_fee_base"] == 240 and anna["registrar_fee_settled"] is True
    # The consultation the case came from is «yes» now, as when a case is opened by hand.
    [consultation] = await rows(
        owner_engine, "SELECT decision, fee_paid FROM consultations WHERE tenant_id = :t", t=tenant
    )
    assert consultation == {"decision": "yes", "fee_paid": True}
    payments = await rows(
        owner_engine,
        "SELECT amount, currency, amount_base, official FROM case_payments WHERE tenant_id = :t "
        "ORDER BY paid_on",
        t=tenant,
    )
    assert [(p["amount"], p["currency"], p["amount_base"], p["official"]) for p in payments] == [
        (200, "BYN", 200, False),
        (150, "EUR", 510, True),
    ]
    [booking] = await rows(
        owner_engine,
        "SELECT b.place, e.login FROM bookings b JOIN vault_entries e ON e.id = b.vault_entry_id "
        "WHERE b.tenant_id = :t",
        t=tenant,
    )
    assert booking["login"] == "booking.mine@example.com"
    attempts = await rows(
        owner_engine,
        "SELECT number, outcome FROM application_attempts WHERE tenant_id = :t ORDER BY number",
        t=tenant,
    )
    assert [(a["number"], a["outcome"]) for a in attempts] == [(1, "expired"), (2, None)]

    # The password is in the vault, encrypted, and comes back only when revealed.
    [entry] = await rows(
        owner_engine,
        "SELECT id, password_encrypted FROM vault_entries WHERE tenant_id = :t "
        "AND login = 'anna.p@example.com'",
        t=tenant,
    )
    assert b"Pa55word" not in entry["password_encrypted"]
    revealed = await ok(await client.post(f"/api/vault/entries/{entry['id']}/reveal"))
    assert revealed["password"] == "Pa55word"

    # One audit event with the counts, nothing from the file.
    [event] = await rows(
        owner_engine,
        "SELECT details::text AS details FROM audit_events WHERE tenant_id = :t "
        "AND action = 'imports.completed'",
        t=tenant,
    )
    for value in ("Петрова", "+375291112233", "Pa55word", "anna.p@example.com"):
        assert value not in event["details"]


async def test_check_writes_nothing_and_lists_the_unknown_list_values(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")

    report = await check(client, {"Клиенты": CLIENTS, "Кейсы": CASES})

    assert report["id"] is None
    assert report["created"]["contacts"] == 3 and report["created"]["cases"] == 3
    assert report["unknown"] == {
        "channel": [{"key": "сарафан", "value": "Сарафан", "rows": 2}],
        "registrar": [{"key": "иван", "value": "Иван", "rows": 2}],
    }
    for table in ("contacts", "visa_cases", "lead_channels", "registrars", "import_runs"):
        assert await count(owner_engine, table, tenant) == 0


async def test_unknown_values_are_added_taken_from_the_list_or_left_empty(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")
    natalia = await ok(
        await client.post(
            "/api/dictionaries/registrars",
            json={"name": "Наталья", "kind": "individual", "currency": "BYN"},
        )
    )

    report = await run(
        client,
        {"Клиенты": CLIENTS, "Кейсы": CASES},
        choices={"channel": {"сарафан": "skip"}, "registrar": {"иван": natalia["id"]}},
    )

    assert "channel" not in report["created"] and "registrar" not in report["created"]
    assert await count(owner_engine, "lead_channels", tenant) == 0
    registrars = await rows(
        owner_engine,
        "SELECT DISTINCT r.name FROM visa_cases v JOIN registrars r ON r.id = v.registrar_id "
        "WHERE v.tenant_id = :t",
        t=tenant,
    )
    assert registrars == [{"name": "Наталья"}]
    assert {"sheet": "Клиенты", "row": 2, "column": "Канал прихода", "reason": "unknown_value"} in (
        report["warnings"]
    )


async def test_a_repeated_or_re_made_file_finds_the_records_it_made(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")
    await mail_system(client)
    await run(client, BOOK)
    counts = {
        table: await count(owner_engine, table, tenant)
        for table in (
            "contacts",
            "contact_notes",
            "consultations",
            "visa_cases",
            "submission_groups",
            "case_payments",
            "case_expenses",
            "waiting_list_entries",
            "application_attempts",
            "bookings",
            "vault_entries",
        )
    }

    again = await run(client, BOOK)
    assert again["created"] == {} and again["updated"] == {}

    # The AI made the file again with other codes: the same records, nothing added.
    renamed = {
        name: [
            [
                f"new-{c}"
                if isinstance(c, str) and "-" in c and c[0].isupper() and c.split("-")[-1].isdigit()
                else c
                for c in row
            ]
            for row in sheet
        ]
        for name, sheet in BOOK.items()
    }
    remade = await run(client, renamed)
    assert remade["skipped_count"] == 0, remade["skipped"]
    assert remade["created"] == {}
    for table, number in counts.items():
        assert await count(owner_engine, table, tenant) == number, table


async def test_people_are_found_by_phone_nick_or_a_unique_name_and_their_card_kept(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")
    for name, extra in (
        ({"last_name": "Петрова", "first_name": "Анна"}, {"phone": "+375291112233"}),
        ({"last_name": "Петров", "first_name": "Игорь"}, {"email": "igor@example.com"}),
        ({"last_name": "Сидоренко", "first_name": "Ольга"}, {"telegram": "olga_s"}),
    ):
        await ok(await client.post("/api/contacts", json=name | extra))

    report = await run(
        client,
        {
            "Клиенты": [
                ["Код клиента", "Фамилия", "Имя", "Телефон", "Telegram", "Email"],
                ["1", "Петрова-Смит", "Анна", "+375291112233", None, "anna@example.com"],
                ["2", "Петров", "Игорь", None, None, "other@example.com"],
                ["3", "Другая", "Фамилия", None, "@olga_s", None],
            ]
        },
    )

    assert report["created"] == {} and report["updated"] == {"contacts": 1}
    contacts = await rows(
        owner_engine, "SELECT full_name, email FROM contacts WHERE tenant_id = :t", t=tenant
    )
    # A file from the AI only fills what the card does not have: names and emails stay.
    assert sorted(contacts, key=lambda c: c["full_name"]) == [
        {"full_name": "Петров Игорь", "email": "igor@example.com"},
        {"full_name": "Петрова Анна", "email": "anna@example.com"},
        {"full_name": "Сидоренко Ольга", "email": None},
    ]


async def test_an_empty_cell_of_a_file_from_the_ai_changes_nothing(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")
    await run(client, {"Клиенты": CLIENTS, "Кейсы": CASES})

    cases = [CASES[0], [*CASES[1][:8], None, *CASES[1][9:]]]
    report = await run(client, {"Клиенты": CLIENTS, "Кейсы": cases})

    assert report["updated"] == {}
    [case] = await rows(
        owner_engine,
        "SELECT vac_number FROM visa_cases WHERE tenant_id = :t AND vac_number IS NOT NULL",
        t=tenant,
    )
    assert case["vac_number"] == "PL123"


async def _exported(owner_engine: AsyncEngine, tenant: uuid.UUID) -> dict[str, dict]:
    """The CRM's own codes of the records, as an export writes them."""
    [anna] = await rows(
        owner_engine,
        "SELECT c.id AS contact, v.id AS case, v.updated_at FROM visa_cases v "
        "JOIN contacts c ON c.id = v.contact_id WHERE v.tenant_id = :t AND v.vac_number = 'PL123'",
        t=tenant,
    )
    payments = await rows(
        owner_engine,
        "SELECT id, amount FROM case_payments WHERE case_id = :c ORDER BY paid_on",
        c=anna["case"],
    )
    return {"anna": anna, "payments": payments}


def _about(at: datetime) -> list[list]:
    return [["Формат", "VisaCRM, версия 1"], ["Выгружено", at]]


def _minsk_now() -> datetime:
    """The export writes its time in the specialist's time zone, without the zone."""
    return datetime.now(ZoneInfo("Europe/Minsk")).replace(tzinfo=None, microsecond=0)


async def test_a_row_of_an_export_is_a_snapshot_of_its_record(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")
    await run(client, {"Клиенты": CLIENTS, "Кейсы": CASES, "Платежи": PAYMENTS})
    crm = await _exported(owner_engine, tenant)
    anna, first = crm["anna"], crm["payments"][0]
    later = _minsk_now() + timedelta(hours=1)

    report = await run(
        client,
        {
            "Кейсы": [
                ["Код кейса", "Код клиента", "Клиент", "Страна визы", "Номер в ВЦ", "Дата подачи"],
                [
                    str(anna["case"]),
                    str(anna["contact"]),
                    "Петрова Анна Сергеевна",
                    "Польша",
                    None,
                    date(2025, 5, 21),
                ],
            ],
            "Платежи": [
                ["Код", "Код кейса", "Дата", "Сумма", "Валюта"],
                [str(first["id"]), str(anna["case"]), date(2025, 1, 15), 250, "BYN"],
            ],
            "О файле": _about(later),
        },
    )

    assert report["skipped_count"] == 0, report["skipped"]
    assert report["created"] == {}
    assert report["updated"] == {"cases": 1, "payments": 1}
    [case] = await rows(
        owner_engine,
        "SELECT vac_number, submission_on FROM visa_cases WHERE id = :c",
        c=anna["case"],
    )
    # The cell left empty in the export clears the field.
    assert case == {"vac_number": None, "submission_on": date(2025, 5, 21)}
    [payment] = await rows(
        owner_engine, "SELECT amount FROM case_payments WHERE id = :p", p=first["id"]
    )
    assert payment["amount"] == 250


async def test_a_record_changed_in_the_crm_after_the_export_is_not_overwritten(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")
    await run(client, {"Клиенты": CLIENTS, "Кейсы": CASES})
    anna = (await _exported(owner_engine, tenant))["anna"]
    earlier = _minsk_now() - timedelta(days=1)

    report = await run(
        client,
        {
            "Кейсы": [
                ["Код кейса", "Код клиента", "Страна визы", "Номер в ВЦ"],
                [str(anna["case"]), str(anna["contact"]), "Польша", "PL999"],
            ],
            "О файле": _about(earlier),
        },
    )

    assert report["skipped"] == [
        {"sheet": "Кейсы", "row": 2, "column": None, "reason": "changed_after_export"}
    ]
    [case] = await rows(
        owner_engine, "SELECT vac_number FROM visa_cases WHERE id = :c", c=anna["case"]
    )
    assert case["vac_number"] == "PL123"


async def test_codes_of_another_workspace_find_nothing(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    other = await make_user("other@example.com")
    await login(client, "other@example.com")
    await run(client, {"Клиенты": CLIENTS, "Кейсы": CASES})
    theirs = (await _exported(owner_engine, other))["anna"]
    client.cookies.clear()
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")

    report = await run(
        client,
        {
            "Платежи": [
                ["Код кейса", "Дата", "Сумма", "Валюта"],
                [str(theirs["case"]), date(2025, 1, 15), 999, "BYN"],
            ],
            "Кейсы": [
                ["Код кейса", "Код клиента", "Страна визы"],
                [str(theirs["case"]), str(theirs["contact"]), "Польша"],
            ],
        },
    )

    # Their client code is no card here and the file has no client sheet: the case row is
    # skipped, and so is the payment of a case the file did not make.
    assert {(s["sheet"], s["reason"]) for s in report["skipped"]} == {
        ("Кейсы", "unknown_code"),
        ("Платежи", "unknown_code"),
    }
    assert await count(owner_engine, "case_payments", tenant) == 0
    assert await count(owner_engine, "case_payments", other) == 0
    [case] = await rows(
        owner_engine, "SELECT vac_number FROM visa_cases WHERE id = :c", c=theirs["case"]
    )
    assert case["vac_number"] == "PL123"


async def test_a_row_whose_client_does_not_match_its_code_is_skipped(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    await login(client, "spec@example.com")

    report = await run(
        client,
        {
            "Клиенты": CLIENTS,
            "Заметки": [
                ["Код клиента", "Клиент", "Текст"],
                ["Подача25-2", "Петров Игорь", "чужая заметка"],
            ],
        },
    )

    assert report["skipped"] == [
        {"sheet": "Заметки", "row": 2, "column": "Клиент", "reason": "client_mismatch"}
    ]
    assert await count(owner_engine, "contact_notes", tenant) == 0


async def test_the_report_names_sheet_row_and_column_and_never_a_value(
    client: AsyncClient, make_user: MakeUser
):
    await make_user("spec@example.com")
    await login(client, "spec@example.com")

    report = await check(
        client,
        {
            "Клиенты": [
                ["Код клиента", "Фамилия", "Имя", "Телефон", "Дата рождения"],
                ["1", "Петрова", "Анна", "8 029 111-22-33", "12.03"],
                ["2", None, None, "+375297778899", None],
            ],
            "Отчёт": [["что угодно"]],
        },
    )

    assert report["skipped"] == [
        {"sheet": "Клиенты", "row": 3, "column": None, "reason": "no_name"}
    ]
    assert report["warnings"] == [
        {"sheet": "Клиенты", "row": 2, "column": "Телефон", "reason": "phone"},
        {"sheet": "Клиенты", "row": 2, "column": "Дата рождения", "reason": "date"},
        {"sheet": "Отчёт", "row": None, "column": None, "reason": "unknown_sheet"},
    ]
    for value in ("Петрова", "029", "12.03", "+375297778899"):
        assert value not in str(report)


async def test_a_sheet_of_a_module_the_workspace_has_not_is_passed_by(
    client: AsyncClient, make_user: MakeUser, owner_engine: AsyncEngine
):
    tenant = await make_user("spec@example.com")
    async with owner_engine.begin() as conn:
        await conn.execute(
            text("UPDATE tenants SET modules = array_remove(modules, 'finances') WHERE id = :t"),
            {"t": tenant},
        )
    await login(client, "spec@example.com")

    report = await run(client, {"Клиенты": CLIENTS, "Кейсы": CASES, "Платежи": PAYMENTS})

    assert {"sheet": "Платежи", "row": None, "column": None, "reason": "module_disabled"} in (
        report["skipped"]
    )
    assert report["created"]["cases"] == 3
    assert await count(owner_engine, "case_finances", tenant) == 0


async def test_only_the_specialist_imports_and_files_are_bounded(
    client: AsyncClient, make_user: MakeUser
):
    tenant = await make_user("spec@example.com")
    await make_user("emp@example.com", role="employee", tenant_id=tenant)
    await login(client, "spec@example.com")

    unreadable = {"file_name": "x.xlsx", "content": base64.b64encode(b"name;phone").decode()}
    response = await client.post(f"{API}/check", json=unreadable)
    assert response.status_code == 422 and response.json()["detail"] == "file_unreadable"

    client.cookies.clear()
    await login(client, "emp@example.com")
    response = await client.post(f"{API}/check", json=body({"Клиенты": CLIENTS}))
    assert response.status_code == 403


async def test_the_history_keeps_each_run_with_its_report(client: AsyncClient, make_user: MakeUser):
    await make_user("spec@example.com")
    await login(client, "spec@example.com")
    report = await run(client, {"Клиенты": CLIENTS})

    [summary] = await ok(await client.get(API))
    assert summary["id"] == report["id"] and summary["kind"] == "workbook"
    assert summary["imported"] == 3
    again = await ok(await client.get(f"{API}/{report['id']}"))
    assert again["created"] == {"contacts": 3, "channel": 1}


async def test_the_template_has_every_sheet_and_loads_back_empty(
    client: AsyncClient, make_user: MakeUser
):
    await make_user("spec@example.com")
    await login(client, "spec@example.com")
    await ok(await client.post("/api/dictionaries/channels", json={"name": "Instagram"}))

    response = await client.get(f"{API}/template")
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == [
        "Клиенты",
        "Заметки",
        "Консультации",
        "Кейсы",
        "Платежи",
        "Расходы",
        "Очередь",
        "Заявки с окном",
        "Брони",
        "Учётные записи",
        "Справочник",
        "О файле",
    ]
    assert [c.value for c in book["Клиенты"][1]][:3] == ["Код клиента", "Фамилия", "Имя"]
    reference = [c.value for column in book["Справочник"].iter_cols() for c in column]
    assert "Instagram" in reference and "Польша" in reference and "многократная" in reference

    content = base64.b64encode(response.content).decode()
    report = await ok(
        await client.post(f"{API}/check", json={"file_name": "t.xlsx", "content": content})
    )
    assert report["rows"] == 0 and report["skipped"] == [] and report["warnings"] == []


async def test_the_prompt_comes_in_both_languages_with_the_format_columns(
    client: AsyncClient, make_user: MakeUser
):
    await make_user("spec@example.com")
    await login(client, "spec@example.com")

    ru = await ok(await client.get(f"{API}/prompt", params={"language": "ru"}))
    en = await ok(await client.get(f"{API}/prompt", params={"language": "en"}))

    assert "Клиенты: Код клиента* | Фамилия | Имя" in ru["text"]
    assert "Clients: Client code* | Last name | First name" in en["text"]
